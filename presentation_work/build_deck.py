import sys,json,math
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'presentation_work/python_lib'))
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from PIL import Image

W=13.333333;H=7.5
NAVY='173D6A';BLUE='0072BC';CYAN='45ADC4';PALE='EAF2FA';INK='22384B';MUTED='526778';LINE='C6D8E8';ORANGE='C77823';GREEN='21815C';WHITE='FFFFFF'
prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
prs.core_properties.title='IPsec Sentinel | SIH 2026 | PS 26160'
prs.core_properties.subject='Evidence-driven technical presentation for NTRO'
prs.core_properties.author='IPsec Sentinel'
logo=ROOT/'presentation_work/assets/image2.png'
slides=[];notes=[];facts=[]
def rect(s,x,y,w,h,fill=PALE,line=None,rounded=True):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded: q.adjustments[0]=.12
    q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
    if line:q.line.color.rgb=RGBColor.from_string(line);q.line.width=Pt(.8)
    else:q.line.fill.background()
    from pptx.oxml.xmlchemy import OxmlElement
    q._element.spPr.append(OxmlElement('a:effectLst'))
    return q
def text(s,t,x,y,w,h,size=20,color=INK,bold=False,font='Calibri',align=None):
    q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=q.text_frame
    f.clear();f.word_wrap=True
    f.margin_left=f.margin_right=Inches(.02);f.margin_top=f.margin_bottom=0
    for i,line in enumerate(t.split('\n')):
        p=f.paragraphs[0] if i==0 else f.add_paragraph();p.text=line
        p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(2)
        if align is not None:p.alignment=align
    return q
def arrow(s,x1,y1,x2,y2,color=BLUE,head=True):
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));q.line.color.rgb=RGBColor.from_string(color);q.line.width=Pt(1.5)
    from pptx.oxml.xmlchemy import OxmlElement
    if head:
        end=OxmlElement('a:tailEnd');end.set('type','triangle');q.line._get_or_add_ln().append(end)
    return q
def card(s,title,body,x,y,w,h,fill=PALE,accent=BLUE,ts=21,bs=18):
    rect(s,x,y,w,h,fill);rect(s,x,y,.055,h,accent,rounded=False)
    text(s,title,x+.18,y+.16,w-.35,.55,ts,accent,True)
    text(s,body,x+.18,y+.72,w-.35,h-.80,bs,INK)
def pill(s,t,x,y,w,color=BLUE):
    rect(s,x,y,w,.32,color);text(s,t,x+.07,y+.035,w-.14,.26,10.5,WHITE,True)
def band(s,t,y=6.35,color=PALE,tc=NAVY):
    rect(s,.55,y,12.23,.47,color);text(s,t,.74,y+.09,11.85,.32,15,tc)
def new(title,kicker,source):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
    text(s,kicker.upper(),.55,.27,9,.3,11,BLUE,True)
    text(s,title,.55,.73,10.55,.75,29,NAVY,True,'Georgia')
    s.shapes.add_picture(str(logo), Inches(11.27),Inches(.22),width=Inches(1.5))
    rect(s,.55,1.55,.63,.055,CYAN,rounded=False)
    text(s,source,.55,6.94,12.2,.22,9.2,MUTED)
    rect(s,0,7.23,W,.27,BLUE,rounded=False)
    text(s,'SMART INDIA HACKATHON 2026   |   PS 26160   |   NTRO',.55,7.285,10,.15,8.5,WHITE)
    text(s,f'{len(prs.slides):02}',12.17,7.265,.6,.2,10,WHITE,True,align=PP_ALIGN.RIGHT)
    slides.append(s);return s
def addmeta(s,say,key,q,a,claims):
    idx=len(prs.slides);notes.append((idx,say,key,q,a));s.notes_slide.notes_text_frame.text=f'{say}\n\nKey technical point: {key}\nEvaluator question: {q}\nAnswer: {a}'
    for claim,source,status,note in claims:facts.append((idx,claim,source,status,note))

# 1
s=new('IPsec Sentinel','Smart India Hackathon 2026','Evidence baseline: repository audit and reproduced results; see PPT_FACT_CHECK.md.')
title=s.shapes[1];title.height=Inches(1.05)
for p in title.text_frame.paragraphs:p.font.size=Pt(48)
s.shapes[3].top=Inches(1.68)
text(s,'AI-Powered IPsec VPN Protocol Analyzer\n& Security Assessment Framework',.6,1.92,8.6,1.15,28,NAVY)
text(s,'From packet evidence to explainable assessment.',.6,3.28,9,.55,23,BLUE,True)
for x,w,a,b in [( .6,2.1,'PS 26160','Software'),(2.92,3.95,'NTRO','National Technical Research Organisation'),(7.1,3.2,'Theme','Blockchain & Cybersecurity')]:
    card(s,a,b,x,4.25,w,1.5,ts=20,bs=16)
for i,(a,c) in enumerate([('OBSERVE',NAVY),('INFER',BLUE),('ASSESS',CYAN)]):
    x=.6+i*3.55;rect(s,x,6.03,3.05,.52,c);text(s,a,x+.1,6.13,2.85,.3,17,WHITE,True,align=PP_ALIGN.CENTER)
    if i<2:arrow(s,x+3.1,6.29,x+3.48,6.29)
addmeta(s,'IPsec Sentinel addresses PS 26160 for NTRO. Our prototype turns IPsec packet captures into three clearly separated outputs: visible protocol evidence, model estimates, and rule-based security assessment. This presentation follows the implemented code and reproducible artifacts. We will show the controlled dataset, the actual model results, and a real dashboard example, including what the prototype cannot yet determine.','Three distinct evidence categories.','Is this a concept or a working prototype?','The repository contains runnable parsers, four trained artifacts, security rules, APIs, tests, and saved dashboard/report evidence.',[('Project identity and SIH metadata','User brief; reference PPT slide 1','VERIFIED','PS 26160; NTRO; Software; Blockchain & Cybersecurity.'),('Working local prototype','src/App.tsx; server/scapy_analyzer.py; demo/','VERIFIED','No production-deployment claim.')])
# 2
s=new('The bottleneck is interpretation','01 / The problem','Evidence: parser unknown states in server/scapy_analyzer.py; LIMITATIONS.md checked against source.')
text(s,'A capture contains evidence. A defensible security judgement needs context.',.6,1.86,11.9,.6,23,NAVY)
labels=[('PCAP','Raw frames'),('Wireshark','Packet inspection'),('IKE / ESP','Protocol evidence'),('Expert review','Manual interpretation'),('Judgement','Finding + rationale')]
for i,(a,b) in enumerate(labels):
    x=.6+i*2.46;card(s,a,b,x,2.95,2.25,1.53,accent=ORANGE if i==3 else BLUE,ts=20,bs=17)
    if i<4:arrow(s,x+2.28,3.67,x+2.42,3.67)
rect(s,7.98,4.73,2.25,.4,'FFF1DF');text(s,'INTERPRETATION BOTTLENECK',8.05,4.81,2.1,.3,10,ORANGE,True,align=PP_ALIGN.CENTER)
text(s,'Visible IKE transforms',.65,5.48,3.9,.4,20,NAVY,True)
text(s,'Encrypted Child-SA details',4.8,5.48,4.2,.4,20,NAVY,True)
text(s,'Incomplete captures',9.16,5.48,3.8,.4,20,NAVY,True)
band(s,'Missing evidence must remain unknown; ESP ciphertext cannot be read as application plaintext.')
addmeta(s,'The hard part is not opening a PCAP. It is deciding which protocol fields are actually visible, which settings need gateway evidence, and which security conclusions are justified. IKE negotiation can expose transforms, while ESP application content remains encrypted. Incomplete captures add uncertainty. Our aim is to structure that investigation and make its evidence boundaries visible.','Visibility is not the same as certainty.','Can ESP reveal the application or full VPN configuration?','It exposes headers and traffic shape, not plaintext or all hidden Child-SA settings.',[('Manual interpretation workflow','Conceptual workflow; server/scapy_analyzer.py','VERIFIED','No measured time-saving claim.'),('Hidden fields remain unavailable','server/scapy_analyzer.py: sa defaults; src/utils/pcapParser.ts','VERIFIED','No decryption capability claimed.')])
# 3
s=new('One workflow, three explicit kinds of evidence','02 / Our solution','Sources: src/App.tsx; src/utils/securityAuditor.ts; backend/app/ml_inference.py; backend/app/ml_assessment.py.')
text(s,'Traditional',.62,1.98,2,.4,20,MUTED,True)
text(s,'PCAP   →   Expert   →   Manual analysis',3,1.97,9,.5,23,MUTED)
text(s,'IPsec Sentinel',.62,2.78,3,.4,20,NAVY,True)
for i,(a,b,c) in enumerate([('Protocol parsing','Decode visible IKE fields\nand ESP headers.',NAVY),('ML inference','Estimate four crypto labels\nfrom capture features.',BLUE),('Security rules','Apply explicit checks;\nretain source and severity.',CYAN)]):
    card(s,a,b,.6+i*4.17,3.43,3.79,1.89,accent=c,bs=20)
text(s,'OBSERVED',.78,5.66,3.7,.34,13,NAVY,True)
text(s,'ML_INFERRED',4.95,5.66,3.7,.34,13,BLUE,True)
text(s,'RULE-BASED FINDING',9.12,5.66,3.7,.34,13,CYAN,True)
band(s,'Analyst review remains essential for uncertain predictions, missing evidence, and security interpretation.')
addmeta(s,'The implemented workflow automates extraction and produces a consistent assessment. It has three different jobs. Parsing decodes visible fields. Random Forest models estimate four configuration labels. Deterministic rules produce findings and scores from the available evidence. These outputs are displayed separately so that an estimate cannot silently become packet-observed fact. The analyst still resolves incomplete or conflicting evidence.','Parsing, prediction and judgement have separate roles.','Does ML determine the security score?','The observed configuration score is computed by TypeScript rules; ML produces separately tagged advisory findings.',[('Three separate processing roles','src/App.tsx; src/utils/securityAuditor.ts; backend/app/ml_assessment.py','VERIFIED','ML-inferred notes are not the observed score.')])
# 4
s=new('Two analysis lanes, one dashboard','03 / System architecture','Sources: src/utils/scapyClient.ts; server/scapy_analyzer.py; backend/app/{feature_extraction,ml_inference,ml_assessment}.py.')
card(s,'PCAP / PCAPNG','Factory capture or analyst upload',.6,2.02,3.6,1.22,bs=16)
arrow(s,4.25,2.63,4.65,2.63)
card(s,'Local Scapy service','POST :8765/analyze',4.7,2.02,3.55,1.22,accent=NAVY,bs=16)
arrow(s,8.3,2.63,8.72,2.63)
card(s,'React dashboard','App.tsx · reports · evidence',8.78,2.02,3.95,1.22,bs=16)
card(s,'A  /  Packet observations','IKE proposals + ESP metadata\n→ securityAuditor.ts rules',.6,3.95,5.65,1.55,accent=NAVY,bs=20)
card(s,'B  /  Model estimates','extractor.py → 18 features → 4 RF models\n→ ml_assessment.py advisory findings',6.75,3.95,5.98,1.55,accent=CYAN,bs=19)
arrow(s,6.4,3.27,3.4,3.88);arrow(s,6.6,3.27,9.7,3.88)
text(s,'Browser fallback: pcapParser.ts',.62,5.8,5.9,.38,18,MUTED)
text(s,'Optional :8770 API → SQLite + gateway correlation',6.77,5.8,6,.38,17,MUTED)
band(s,'Separate FastAPI interface: /api/ml-analyze. The dashboard’s default upload path is the Scapy service.')
addmeta(s,'The dashboard sends captures to the local Scapy analyzer on port 8765. That service returns decoded observations and invokes the shared feature extractor and four model artifacts. The browser computes its observed scorecard and renders the combined result. A separate local API on 8770 handles persistence and gateway correlation. FastAPI also exposes an ML endpoint. If Scapy is unavailable, browser parsing remains available, but Python model inference is absent.','The diagram matches actual frontend calls.','Is FastAPI the dashboard’s only backend?','No. The default upload uses the Scapy HTTP service; FastAPI is a separate interface.',[('Default frontend path','src/utils/scapyClient.ts; src/App.tsx','VERIFIED','8765 analyzer; 8770 registration/telemetry.'),('FastAPI alternative','backend/app/main.py','VERIFIED','Not represented as the default UI path.'),('Browser fallback','src/App.tsx; src/utils/pcapParser.ts','VERIFIED','Does not execute Python model artifacts.')])
# 5
s=new('Ground truth begins in a controlled VPN lab','04 / PCAP factory','Sources: pcap_factory/{factory.py,experiments.json,metadata/}; feature_extractor/dataset.csv; evidence/dataset_audit.json.')
for i,(a,b) in enumerate([('157','successful captures'),('24','crypto combinations'),('113 / 44','ICMP / TCP captures')]):
    x=.6+i*4.17;rect(s,x,1.95,3.79,1.08,PALE);text(s,a,x+.2,2.06,3.4,.5,31,BLUE,True);text(s,b,x+.2,2.63,3.4,.28,16,MUTED)
steps=[('Define','experiments.json'),('Configure','2 strongSwan peers'),('Capture + traffic','tcpdump + ping/TCP'),('Verify','SA state + PFS rekey'),('Label','PCAP + metadata')]
for i,(a,b) in enumerate(steps):
    x=.6+i*2.46;card(s,a,b,x,3.55,2.25,1.45,ts=18,bs=15.5)
    if i<4:arrow(s,x+2.28,4.27,x+2.42,4.27)
text(s,'44 definitions × 5 intended repetitions',.65,5.43,6.2,.43,22,NAVY,True)
text(s,'157 SUCCESS + 54 FAILED metadata records',.65,5.95,6.3,.35,18,MUTED)
text(s,'33 captured crypto/traffic scenarios\nAll successes: IKEv2 + tunnel mode',7.05,5.42,5.6,.87,20,NAVY)
band(s,'Labels are parsed from filenames; saved negotiation metadata provides the separate validation trail.',6.45)
addmeta(s,'The factory configures two strongSwan peers, begins packet capture, generates traffic, and compares negotiated settings with the experiment. PFS-enabled runs require explicit rekey evidence. We counted 157 successful captures covering all 24 crypto combinations, across 33 crypto-and-traffic scenarios. The repository also retains 54 failure records. Forty-four definitions with five repetitions describe the intended matrix, not a completed campaign. CSV labels come from filenames, with negotiation metadata checked separately.','Configuration labels have a saved negotiation trail.','Is this a production dataset or synthetic packet generator?','These are controlled strongSwan lab captures with negotiation metadata. The browser’s separate testbed export is synthetic.',[('157 captures; 24 combinations; 33 scenarios; 113 ICMP/44 TCP','presentation_work/evidence/dataset_audit.json; dataset.csv','VERIFIED','No successful UDP captures.'),('44 definitions × 5; 157 SUCCESS / 54 FAILED','pcap_factory/experiments.json; metadata/','VERIFIED','Do not infer 220 completed attempts; records may reflect retries.'),('Verification and labels','factory.py validate_negotiation/perform_child_rekey; build_dataset.py','VERIFIED','Filename label extraction does not itself read/validate negotiation metadata.')])
# 6
s=new('Extract features without decrypting ESP','05 / From packets to features','Sources: feature_extractor/extractor.py; backend/app/feature_extraction.py; server/scapy_analyzer.py; src/utils/aiClassifier.ts.')
card(s,'Visible protocol evidence','IKE version and exchange headers\nSA transform IDs and key length\nESP SPI and sequence numbers',.6,1.96,5.8,1.91,accent=NAVY,bs=20)
card(s,'Encrypted-flow observations','Packet sizes, byte counts and duration\nTiming, symmetry and entropy in UI\nApplication plaintext remains hidden',6.7,1.96,6.03,1.91,accent=CYAN,bs=20)
text(s,'EXACT INPUTS TO THE FOUR CRYPTO MODELS  /  18 FEATURES',.63,4.13,12,.35,12,BLUE,True)
groups=[['packet_count','total_bytes','avg_packet_size','min_packet_size','max_packet_size','capture_duration_seconds'],['udp_packet_count','udp_500_count','ike_packet_count','esp_packet_count','create_child_sa_count','informational_count'],['ike_request_count','ike_response_count','ike_bytes','ike_avg_packet_size','esp_bytes','esp_avg_packet_size']]
for i,g in enumerate(groups):
    rect(s,.6+i*4.17,4.61,3.79,1.62,PALE);box=text(s,'\n'.join(g),.78+i*4.17,4.73,3.45,1.5,14.5,NAVY,font='Consolas')
    for p in box.text_frame.paragraphs:p.space_after=Pt(0)
band(s,'Timing, symmetry and entropy are separate UI flow statistics; they are not inputs to these 4 crypto models.')
addmeta(s,'We distinguish decoded protocol fields from measured packet statistics. The protocol parser exposes visible headers and transforms. The crypto models use exactly these eighteen numeric features, including packet counts, sizes, duration and IKE/ESP totals. The dashboard also measures timing, symmetry and entropy for a separate workload heuristic. Those are not the four models’ inputs. The encrypted ESP payload is never treated as readable application content.','Use the actual trained feature schema.','Do entropy or packet sizes prove the cipher?','No. They support statistical estimates at most; visible transforms or correlated gateway evidence are stronger confirmation.',[('18 exact trained features','backend/app/feature_extraction.py ML_FEATURE_NAMES; model_metadata.joblib','VERIFIED','No invented inputs; schema validation checks presence and inference orders columns.'),('UI workload feature separation','server/scapy_analyzer.py; src/utils/aiClassifier.ts','VERIFIED','These statistics are not crypto-model inputs.'),('Protocol support differs in ML extractor','feature_extractor/extractor.py','PARTIAL','Native ESP extraction is IPv4-specific; UDP/4500 fallback is simplistic; see audit.')])
# 7
s=new('Four models. Four configuration estimates.','06 / ML pipeline','Sources: train_final_models.py; make_split.py; model_metadata.joblib; backend/app/ml_inference.py.')
steps=[('PCAP','Scapy extraction'),('18 features','Select + order columns'),('4 forests','300 trees per model'),('Prediction','Class + probability')]
for i,(a,b) in enumerate(steps):
    x=.6+i*3.105;card(s,a,b,x,2.04,2.82,1.43,ts=20,bs=16)
    if i<3:arrow(s,x+2.86,2.73,x+3.07,2.73)
rows=[('Encryption','AES128  ·  AES256'),('Hash / integrity','SHA256  ·  SHA384'),('DH group','DH14  ·  DH15'),('PFS class','NOPFS  ·  PFS14  ·  PFS15')]
for i,(a,b) in enumerate(rows):
    y=3.93+i*.51;rect(s,.6,y,7.25,.45,PALE if i%2==0 else 'F5F8FC',rounded=False);text(s,a,.76,y+.075,2.3,.31,17,NAVY,True);text(s,b,3.2,y+.075,4.45,.31,17,INK)
card(s,'Training contract','126 rows · repetitions 1–4\nBalanced class weights · seed 42\nNo scaler in bundled training',8.18,3.93,4.55,1.98,bs=18)
band(s,'No mode, IKE-version or traffic-type target. Probabilities are uncalibrated; unseen classes are not rejected.')
addmeta(s,'The trained component is deliberately specific. There are four Random Forest classifiers, each with three hundred trees, predicting encryption, integrity, DH group and PFS class. Training uses 126 rows from repetitions one through four, with balanced class weights and a fixed seed. Inference selects the same eighteen columns in the saved order. There is no tunnel-mode model or application classifier here, and predicted-class probability is not measured accuracy.','Closed-set crypto estimates, not arbitrary protocol discovery.','What happens for an unseen cipher or transport mode?','The crypto model still selects among known labels; no OOD rejection exists. Mode is not a target and stays undetermined without other evidence.',[('4 RF models; 300 trees; 126 training rows; seed 42','feature_extractor/ml/train_final_models.py; model_metadata.joblib','VERIFIED','No retraining performed for this deck.'),('Exact classes','presentation_work/evidence/model_evaluation.json; joblib classes_','VERIFIED','2/2/2/3 classes across separate targets.'),('Unknown-class and mode limitations','backend/app/ml_inference.py; model artifacts','PARTIAL','No OOD detector, calibration or mode predictor.')])
# 8
s=new('Measured results show both capability and limits','07 / Model validation','Reproduced from bundled joblib files + test.csv; full metrics and confusion matrices: presentation_work/evidence/model_evaluation.json.')
text(s,'Held-out accuracy  /  Repetition 5  /  31 captures',.6,1.92,9.1,.45,23,NAVY,True)
vals=[('Cipher',19,61.29),('Hash',31,100),('DH group',29,93.55),('PFS',30,96.77)]
for i,(label,n,pct) in enumerate(vals):
    y=2.74+i*.68;text(s,label,.62,y,1.45,.38,20,NAVY,True);rect(s,2.16,y+.02,4.5,.32,PALE,rounded=False);rect(s,2.16,y+.02,4.5*pct/100,.32,ORANGE if i==0 else BLUE,rounded=False);text(s,f'{pct:.2f}%  ({n}/31)',6.87,y,2.5,.42,20,NAVY)
text(s,'CIPHER CONFUSION MATRIX',9.36,2.29,3.5,.31,12,BLUE,True)
text(s,'Predicted',10.18,2.8,2.2,.3,14,MUTED,align=PP_ALIGN.CENTER)
text(s,'AES128     AES256',10.08,3.17,2.4,.34,14,NAVY)
for row,lab in enumerate(['AES128','AES256']):
    text(s,lab,9.28,3.71+row*.65,.92,.3,13,NAVY)
    for col,v in enumerate([[6,7],[5,13]][row]):
        x=10.32+col*1.02;y=3.59+row*.65;rect(s,x,y,.94,.58,PALE if row!=col else 'CDE4F4',rounded=False);text(s,str(v),x,y+.1,.94,.4,21,NAVY,True,align=PP_ALIGN.CENTER)
text(s,'Rows = actual labels',9.56,5.08,3,.33,13,MUTED)
text(s,'Validation boundary',.65,5.74,3,.4,19,ORANGE,True)
text(s,'Familiar lab configurations recur across the split.\nNo external-network validation or calibrated confidence yet.',3.55,5.64,9.13,.78,19,INK)
band(s,'Software checks: server 55/55 and FastAPI 22/22 pass; frontend 17/19 pass (2 report-text expectations fail).',6.48)
addmeta(s,'We reran inference using the saved models and the 31-row repetition-five test set. Cipher accuracy is 61.29 percent; hash is 100 percent; DH is 93.55 percent; PFS is 96.77 percent. The cipher confusion matrix makes its errors visible. These are small controlled-lab results with recurring configurations across train and test, not proof of generalization. Existing test suites also reveal two failing frontend report-text expectations, recorded in the fact check.','Separate reproducibility from external validity.','Why is hash perfect but cipher weaker?','These figures reflect this small split and lab capture patterns; the repository does not establish a causal or cross-network explanation.',[('Accuracy and cipher confusion matrix','presentation_work/evidence/model_evaluation.json; feature_extractor/ml/test.csv','VERIFIED','Evaluation of existing artifacts; n=31; no held-out-organization claim.'),('Test suite counts','presentation_work/evidence/{typescript_tests,python_tests,backend_tests}.txt','VERIFIED','17/19; 55/55; 22/22. Two report-string assertions fail.'),('Robust generalization','make_split.py; dataset.csv','FUTURE','Need configuration/device/network-disjoint validation and calibration.')])
# 9
s=new('Security scoring keeps risk and evidence separate','08 / Security assessment engine','Sources: src/utils/securityAuditor.ts; backend/app/ml_assessment.py; evidence/sample_assessment.json.')
steps=[('Observed setting','AES-CBC'),('Implemented rule','Prefer AEAD'),('Finding + evidence','Low · penalty 10')]
for i,(a,b) in enumerate(steps):
    x=.6+i*4.17;card(s,a,b,x,2.03,3.79,1.37,ts=20,bs=21)
    if i<2:arrow(s,x+3.84,2.72,x+4.12,2.72)
for i,(big,small,c) in enumerate([('10 / 100','Observed rule penalty',ORANGE),('60%','Evidence coverage',BLUE),('54 / 100','Evidence-adjusted security score',NAVY)]):
    x=.6+i*4.17;rect(s,x,3.94,3.79,1.28,PALE);text(s,big,x+.19,4.08,3.4,.52,30,c,True);text(s,small,x+.19,4.74,3.43,.34,16,MUTED)
text(s,'Sample calculation: (100 − 10) × 0.60 = 54',.65,5.62,7.2,.4,22,NAVY,True)
text(s,'PFS • replay • lifetime: unknown',8.2,5.65,4.5,.4,18,MUTED)
band(s,'CBC is a hardening note, not proof of an exploit. ML-inferred findings and standards labels need confirmation.')
addmeta(s,'The sample’s visible AES-CBC setting triggers the implemented ten-point hardening penalty. Evidence coverage is sixty percent because PFS, replay configuration and lifetime are unavailable. The evidence-adjusted security score is therefore ninety times point six, or fifty-four. These are local rule outputs, not attack probabilities. Separately, the model predicts no PFS and creates an inferred advisory; that must not be promoted to a confirmed configuration finding.','A low score can reflect incomplete evidence, not just observed risk.','Does a score of 54 mean the tunnel has been compromised?','No. It combines a policy penalty and evidence coverage. Neither the score nor CBC observation establishes exploitation.',[('10 risk; 60% coverage; 54 security score','evidence/sample_assessment.json; src/utils/securityAuditor.ts','VERIFIED','Score formula reproduced on sample.'),('CBC rule and ML advisory','securityAuditor.ts; backend/app/ml_assessment.py','VERIFIED','Rule output, not a confirmed vulnerability/exploit.'),('NIST/standards mapping','securityAuditor.ts','PARTIAL','Baseline rule labels; sample NIST is unverified; no certification.')])
# 10
s=new('A real dashboard exposes the evidence boundary','09 / Product evidence','Real project screenshot (cropped): demo/dark-analysis-preview.png. Controlled ICMP sample; external annotations.')
pic=s.shapes.add_picture(str(ROOT/'demo/dark-analysis-preview.png'),Inches(.6),Inches(1.85),width=Inches(8.42))
pic.crop_top=.10;pic.height=Inches(4.73625)
callouts=[('01  Capture identity','Uploaded file + 46 frames'),('02  Observed settings','AES-CBC · IKEv2 · DH14'),('03  Model confidence','93% mean class probability'),('04  Finding + gaps','10 risk · 60% evidence')]
for i,(a,b) in enumerate(callouts):
    y=1.93+i*1.05;card(s,a,b,9.3,y,3.43,.93,ts=16,bs=14)
    # Tighter body positioning for compact callouts.
    box=s.shapes[-1];box.top=Inches(y+.52);box.height=Inches(.32)
rect(s,9.3,6.19,3.43,.63,'FFF1DF');text(s,'VoIP is a heuristic guess.\nThis sample’s label is ICMP.',9.45,6.28,3.13,.5,13.5,ORANGE,True)
addmeta(s,'This is the repository’s actual dashboard screenshot for the controlled ICMP capture. It shows the uploaded file, visible cipher and DH information, explicit unknown mode, risk and coverage, and the model-confidence summary. One useful limitation is visible: the workload heuristic ranks this ICMP example as VoIP. We are not presenting that label as successful application detection. The ninety-three percent card averages crypto-model probabilities; it does not validate the workload guess.','Real UI evidence includes real limitations.','Why does the dashboard say VoIP for an ICMP test?','Its workload function uses hand-written size/timing scores and is not validated as an application classifier.',[('Real screenshot','demo/dark-analysis-preview.png; demo/README.md','VERIFIED','Screenshot reused without fabricated UI values.'),('93% crypto summary','sample_analysis.json; src/utils/assessmentReport.ts','VERIFIED','Mean 93.25% rounded to 93%; not accuracy.'),('Workload classification limitation','aiClassifier.ts; sample_assessment.json; factory metadata','PARTIAL','Controlled ICMP sample ranked VoIP by heuristic; not a trained workload result.')])
# 11
s=new('One capture, one traceable result','10 / End-to-end demonstration','Reproduced: evidence/sample_analysis.json + sample_assessment.json. Source PCAP and matching factory metadata are preserved.')
pill(s,'INPUT',.62,1.9,.82)
text(s,'AES128_SHA256_DH14_NOPFS_ICMP_rep01.pcap',1.65,1.9,10.9,.43,21,NAVY,True)
text(s,'46 frames  →  6 IKE + 40 ESP  →  extraction + inference + assessment',.63,2.53,12,.48,22,BLUE)
card(s,'Packet-observed','IKEv2\nAES-CBC · 128-bit\nSHA2-256 integrity\nDH Group 14',.6,3.3,3.79,2.5,accent=NAVY,bs=21)
card(s,'Model estimates','AES128      73%\nSHA256    100%\nDH14         100%\nNOPFS      100%',4.78,3.3,3.79,2.5,accent=BLUE,bs=21)
card(s,'Rule outputs','CBC hardening note\nRisk penalty: 10 / 100\nEvidence coverage: 60%\nSecurity score: 54 / 100',8.96,3.3,3.77,2.5,accent=CYAN,bs=19)
band(s,'Observed PFS and mode stay unknown. NOPFS is a model estimate; 100% model probability is not certainty.')
addmeta(s,'Here is the exact demonstration file. The analyzer reads forty-six frames, including six IKE and forty ESP packets. The parser sees IKEv2, AES-CBC with a 128-bit key, SHA2-256 integrity and DH14. The models return the four predictions shown, while the observed rules produce the CBC note and the scores. The factory knows the controlled configuration, but the uploaded-capture path still correctly leaves hidden PFS and mode unknown.','One sample shows all three evidence categories together.','Is this demo an independent model test?','No. Repetition one is in the training partition. This is pipeline demonstration; slide 8 uses repetition-five held-out rows.',[('46 frames; 6 IKE; 40 ESP; observed values','evidence/sample_analysis.json; source PCAP','VERIFIED','Header/proposal observations; no Child-SA plaintext claim.'),('73/100/100/100 probabilities','evidence/sample_analysis.json','VERIFIED','Training-set example, explicitly not validation.'),('Sample rule outputs','evidence/sample_assessment.json','VERIFIED','Actual TypeScript functions rerun; PFS/mode remain unknown.')])
# 12
s=new('Local deployment is feasible; hardening remains','11 / Feasibility and deployment','Sources: package.json; server/{api_server,telemetry,live_capture,repository}.py; Dockerfile; docker-compose.yml.')
steps=[('Analyst browser','React 19 + Vite'),('Local analyzer','Python + Scapy + sklearn'),('Local API','SQLite + gateway evidence')]
for i,(a,b) in enumerate(steps):
    x=.6+i*4.17;card(s,a,b,x,1.97,3.79,1.34,ts=20,bs=17)
    if i==0:arrow(s,x+3.84,2.62,x+4.12,2.62)
arrow(s,2.59,3.32,2.59,3.54,head=False)
arrow(s,2.59,3.54,10.94,3.54,head=False)
arrow(s,10.94,3.54,10.94,3.32)
card(s,'Implemented locally','PCAP upload + browser fallback\nFour bundled model artifacts\nMarkdown/PDF reports + SQLite\nOptional strongSwan SPI correlation',.6,3.83,5.88,2.23,accent=GREEN,bs=20)
card(s,'Before operational deployment','Validate on unseen devices/networks\nCalibrate confidence + add rejection\nComplete container ML packaging\nDeployment security + load testing',6.82,3.83,5.91,2.23,accent=ORANGE,bs=20)
band(s,'Bounded live-capture agent code exists. NIC capture and production deployment were not demonstrated in this audit.')
addmeta(s,'The practical deployment today is local: a React interface, a Python Scapy analyzer, and an API with SQLite persistence and optional gateway evidence. The repository includes bounded live-capture agent code and strongSwan telemetry support, but this audit did not capture from a live interface. Docker scaffolding is present, yet its backend image omits shared ML modules and artifacts. Operational use therefore requires packaging work, deployment hardening, and external validation.','Runnable local components are distinct from proven production deployment.','Can we deploy the Docker configuration unchanged with ML?','The current backend image copies server code only, so ML dependencies in backend and feature_extractor must be packaged first.',[('Actual stack and local services','package.json; server/scapy_analyzer.py; server/api_server.py','VERIFIED','Local prototype; optional external narrative service not required for deterministic report.'),('Gateway/live capture','server/telemetry.py; correlation.py; live_capture.py; agent','PARTIAL','Implemented code and tests; no NIC or new gateway demo in this audit.'),('Docker ML packaging','Dockerfile; docker-compose.yml','PARTIAL','Missing backend and feature_extractor copies; no container runtime claim.'),('Operational hardening','Presentation roadmap grounded in audit gaps','FUTURE','Not completed capabilities.')])
# 13
s=new('The value is a consistent, reviewable investigation','12 / Impact and PS alignment','Evidence: parser outputs, dataset metadata, assessmentReport.ts and saved report exports. No time-saving benchmark is claimed.')
text(s,'Manual investigation',.64,1.97,5.8,.45,24,MUTED,True)
text(s,'IPsec Sentinel prototype',6.94,1.97,5.8,.45,24,NAVY,True)
items=[('Inspect packets individually','Structured protocol observations'),('Reconstruct the configuration','Model estimates with explicit provenance'),('Write the assessment by hand','Rule findings + evidence gaps + report')]
for i,(a,b) in enumerate(items):
    y=2.78+i*.92;rect(s,.6,y,5.43,.67,'F3F6F9');text(s,a,.79,y+.16,5.02,.38,20,MUTED);arrow(s,6.15,y+.34,6.68,y+.34);rect(s,6.8,y,5.93,.67,PALE);text(s,b,6.99,y+.16,5.52,.38,19,NAVY,True)
text(s,'Intended users',.65,5.88,2.4,.4,19,BLUE,True)
text(s,'SOC analysts  •  Government / defence teams  •  Network auditors',3.1,5.9,9.6,.4,19,NAVY)
band(s,'Demonstrated: repeatable processing and report output. Analyst time savings and field impact remain unmeasured.')
addmeta(s,'The prototype’s demonstrated value is standardization. Instead of reconstructing every result manually, an analyst gets structured observations, explicitly labeled model estimates, rule findings, and a reusable report. This maps directly to the problem statement’s analysis and assessment workflow. SOC teams, government or defence network teams, and auditors are intended users. We have not benchmarked analyst time savings or measured deployment impact, so those remain evaluation goals rather than claimed outcomes.','The benefit is grounded in implemented outputs.','How much analyst time does this save?','No controlled time study exists yet. The repository proves processing and report generation, not a speedup percentage.',[('Structured analysis and reports','src/utils/assessmentReport.ts; demo/sample-technical-report.pdf','VERIFIED','No quantified productivity or adoption claim.'),('Intended audiences','User brief; prototype workflow','FUTURE','Target users, not deployed customers.')])
# 14
s=new('From encrypted IPsec packet captures\nto explainable security assessment.','IPsec Sentinel / Final takeaway','Evidence pack: PPT_FACT_CHECK.md • SPEAKER_NOTES.md • presentation_work/evidence/ • rendered_slides/')
text(s,'Observe what is visible. Label what is inferred. Show what is unknown.',.62,2.45,12.03,.87,25,BLUE,True)
steps=[('Packet evidence','IKE + ESP metadata'),('Bounded ML role','4 crypto estimates'),('Explainable output','Findings + provenance')]
for i,(a,b) in enumerate(steps):
    x=.6+i*4.17;card(s,a,b,x,3.88,3.79,1.48,accent=[NAVY,BLUE,CYAN][i],ts=22,bs=19)
    if i<2:arrow(s,x+3.84,4.62,x+4.12,4.62)
text(s,'SIH 2026   /   PS 26160   /   NTRO',.63,6.07,11.95,.55,24,NAVY,True,align=PP_ALIGN.CENTER)
addmeta(s,'Our closing claim is deliberately precise: IPsec Sentinel turns encrypted IPsec captures into an explainable assessment of the evidence that is available. The prototype has a controlled dataset, reproducible model evaluation, a working analysis flow and real report output. Its next step is broader validation and operational hardening. The principle remains the same: observe what is visible, label what is inferred, and show what is still unknown.','Credibility comes from visible evidence and bounded claims.','What is the next technical milestone?','Configuration- and network-disjoint validation, calibrated/OOD-aware inference, and verified deployment packaging.',[('Closing capability summary','Repository audit; slides 4–12 evidence','VERIFIED','Prototype scope, not autonomous or complete protocol discovery.')])

# Final slide has a two-line heading; give it safe space without altering the body.
slides[-1].shapes[1].height=Inches(1.22)
slides[-1].shapes[3].top=Inches(2.08)

# Write required notes, fact check and editable deck.
speaker=['# Speaker notes','Each script is intended for approximately 20–40 seconds. Paths are relative to the repository.','']
for n,say,key,q,a in notes:
    speaker += [f'## Slide {n}',say,'',f'**Key technical point:** {key}',f'**Likely question:** {q}',f'**Answer:** {a}','']
(ROOT/'SPEAKER_NOTES.md').write_text('\n'.join(speaker),encoding='utf8')
fact=['# Presentation fact check','Audit date: 2026-09-28. Read `presentation_work/REPOSITORY_AUDIT.md` for the complete capability matrix and known discrepancies. Metrics were recomputed from existing artifacts; models and application code were not changed.','', '| Slide | Claim | Source file / module | Status | Notes |','|---|---|---|---|---|']
for row in facts:fact.append('| '+' | '.join(str(v).replace('|','/') for v in row)+' |')
fact+=['','## Reproduction','- `python presentation_work/audit.py` → dataset audit, capture inventory, evaluation metrics and sample Scapy output. Requires existing model dependencies plus the local presentation libraries.','- `node_modules/.bin/tsx presentation_work/verify_sample.ts` → actual TypeScript scorecard and workload heuristic.','- `python presentation_work/build_deck.py` → editable PPTX, this file and speaker notes.','- `powershell -File presentation_work/render.ps1` → actual PowerPoint slide images and text-bound checks.','- Existing suites: `npm test`; `python -m unittest discover -s server -p "test_*.py"`; `python -m pytest backend/tests -q`. Logs preserved under `presentation_work/evidence/`.','', '## Deliberate exclusions','No production accuracy, unseen-class detection, trained mode prediction, trained application recognition, full IKEv1 decoding, verified standards compliance, exploit detection, production deployment, or measured time savings are claimed. Reference image1.png (2022) is excluded; only image2.png (2026) is used.','', '## Existing artifact caveats','The dashboard screenshot includes its original heuristic VoIP label for an ICMP sample, standards labels, and a confidence summary. Slide 10 explicitly qualifies those values; none are used as proof of application identity or certification. The sample report’s prose is not used as protocol ground truth. Slide 11 uses a training-partition example for demonstration, while slide 8 uses held-out repetition 5.']
(ROOT/'PPT_FACT_CHECK.md').write_text('\n'.join(fact),encoding='utf8')
out=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'SIH2026_IPsec_Sentinel_Final.pptx';prs.save(out)
print(f'Saved {len(prs.slides)} slides to {out}')
