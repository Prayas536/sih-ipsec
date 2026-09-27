import sys, json, hashlib, zipfile, collections, contextlib, io, sqlite3, re
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT),str(ROOT/'presentation_work/python_lib')]
OUT=ROOT/'presentation_work/evidence'
OUT.mkdir(exist_ok=True)
def save(name,value): (OUT/name).write_text(json.dumps(value,indent=2,default=str),encoding='utf8')
files=[]; source_index=[]
for p in ROOT.rglob('*'):
    if not p.is_file() or any(x in p.parts for x in ['.git','node_modules','presentation_work','__pycache__','dist']): continue
    if p.name=='.env': continue
    rel=str(p.relative_to(ROOT)); files.append({'file':rel,'bytes':p.stat().st_size})
    if p.suffix in ['.py','.ts','.tsx','.md','.json','.yml','.txt','.css'] and 'metadata' not in p.parts and 'output' not in p.parts:
        t=p.read_text(encoding='utf8',errors='replace')
        source_index.append({'file':rel,'lines':len(t.splitlines()),'symbols':re.findall(r'^(?:export )?(?:async )?(?:def |class |function |const )(\w+)',t,re.M),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
save('repository_inventory.json',files);save('source_index.json',source_index)
ex=json.loads((ROOT/'pcap_factory/experiments.json').read_text())
meta=[(p,json.loads(p.read_text())) for p in (ROOT/'pcap_factory/metadata').glob('*.json')]
caps=list((ROOT/'pcap_factory/captures').glob('*.pcap'))
import pandas as pd, joblib
from sklearn.metrics import classification_report, confusion_matrix, accuracy_score
df=pd.read_csv(ROOT/'feature_extractor/dataset.csv');train=pd.read_csv(ROOT/'feature_extractor/ml/train.csv');test=pd.read_csv(ROOT/'feature_extractor/ml/test.csv')
summary={'inventory_files':len(files),'source_files':len(source_index),'captures':len(caps),'feature_json':len(list((ROOT/'feature_extractor/output').glob('*.json'))),'experiments_defined':len(ex['experiments']),'repetitions':ex['settings']['repetitions'],'metadata_records':len(meta),'metadata_status':dict(collections.Counter(m.get('status') for p,m in meta)),'dataset_rows':len(df),'train_rows':len(train),'test_rows':len(test),'traffic':df.traffic.value_counts().to_dict(),'captured_scenarios':df[['encryption','hash','dh_group','pfs_group','traffic']].drop_duplicates().shape[0],'crypto_combinations':df[['encryption','hash','dh_group','pfs_group']].drop_duplicates().shape[0],'metadata_modes':dict(collections.Counter(m.get('mode') for p,m in meta if m.get('status')=='SUCCESS')),'metadata_valid_success':sum(m.get('validation',{}).get('valid') is True for p,m in meta if m.get('status')=='SUCCESS')}
evaluations={}
for target in ['encryption','hash','dh_group','pfs_group']:
    pack=joblib.load(ROOT/f'feature_extractor/ml/models/{target}_model.joblib');model=pack['model']; feats=pack['features']; pred=model.predict(test[feats])
    evaluations[target]={'accuracy':accuracy_score(test[target],pred),'correct':int((pred==test[target]).sum()),'n':len(test),'report':classification_report(test[target],pred,output_dict=True,zero_division=0),'classes':list(model.classes_),'confusion_matrix':confusion_matrix(test[target],pred,labels=model.classes_).tolist(),'features':feats,'model':type(model).__name__,'n_estimators':model.n_estimators}
summary['train_test_exact_feature_overlap']=sum(tuple(row) in set(map(tuple,train[feats].values)) for row in test[feats].values)
summary['model_metadata']=joblib.load(ROOT/'feature_extractor/ml/models/model_metadata.joblib')
save('dataset_audit.json',summary);save('model_evaluation.json',evaluations)
from server.scapy_analyzer import analyze
sample=ROOT/'pcap_factory/captures/AES128_SHA256_DH14_NOPFS_ICMP_rep01.pcap'
with contextlib.redirect_stdout(io.StringIO()): result=analyze(sample.read_bytes(),sample.name)
save('sample_analysis.json',result)
# Read every capture's packet framing; aggregate without exposing packet payloads.
from scapy.all import RawPcapReader
capture_inventory=[]
for p in caps:
    try:
        with RawPcapReader(str(p)) as r: sizes=[len(b) for b,m in r]
        capture_inventory.append({'file':p.name,'packets':len(sizes),'bytes':sum(sizes),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    except Exception as e: capture_inventory.append({'file':p.name,'error':str(e)})
save('capture_inventory.json',capture_inventory)
dbs=[]
for p in ROOT.rglob('*.sqlite3'):
    if 'presentation_work' in p.parts: continue
    c=sqlite3.connect(p.as_uri()+'?mode=ro',uri=True)
    tables=[r[0] for r in c.execute("select name from sqlite_master where type='table'")]
    dbs.append({'file':str(p.relative_to(ROOT)),'tables':{t:c.execute('select count(*) from "'+t.replace('"','""')+'"').fetchone()[0] for t in tables}});c.close()
save('database_inventory.json',dbs)
from pptx import Presentation
deck=Presentation(ROOT/'SIH2026_IPsec_VPN_Analyzer.pptx')
slides=[]
for i,s in enumerate(deck.slides,1):
    slides.append({'slide':i,'texts':[sh.text for sh in s.shapes if sh.has_text_frame]})
save('reference_deck.json',{'width':deck.slide_width,'height':deck.slide_height,'slides':slides})
with zipfile.ZipFile(ROOT/'SIH2026_IPsec_VPN_Analyzer.pptx') as z:
    for n in z.namelist():
        if n.startswith('ppt/media/'):
            (ROOT/'presentation_work/assets'/Path(n).name).write_bytes(z.read(n))
print(json.dumps(summary,indent=2,default=str))
print('EVALUATION', {k:(v['correct'],v['n'],v['accuracy']) for k,v in evaluations.items()})
print('SAMPLE',result['mlPredictions'],result['mlWarning'])
